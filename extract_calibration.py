#!/usr/bin/env python3
"""Build audited JSON parameters from the upstream DP4+/STS spreadsheets."""
import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import ast


def main():
    import openpyxl
    import msoffcrypto
    p = argparse.ArgumentParser()
    p.add_argument('asset_directory')
    p.add_argument('--dp4-source',required=True)
    args = p.parse_args()
    asset = Path(args.asset_directory)
    root = Path(__file__).resolve().parent
    (root/'upstream').mkdir(exist_ok=True)
    with (asset/'custom-dp4plus.xlsx').open('rb') as stream:
        office = msoffcrypto.OfficeFile(stream)
        # Excel's documented transparent/default encryption, not password guessing.
        office.load_key(password='VelvetSweatshop',verify_password=True)
        readable = io.BytesIO()
        office.decrypt(readable)
    (root/'upstream/DP4plus.xlsx').write_bytes(readable.getvalue())
    shutil.copy2(asset/'STS.xlsx',root/'upstream/STS.xlsx')
    w = openpyxl.load_workbook(readable,data_only=False)
    s, dp = w['Statistics'], {}
    for col in range(3,27):
        def val(row):
            return s.cell(row,col).value
        key = '|'.join([val(1),val(2),val(3)])
        dp[key] = {'functional':val(1),'solvation':val(2),'basis':val(3),'TMS':{'C':val(5),'H':val(6)},
          'unscaled':{'C_sp2':[val(8),val(9),val(10)],'C_sp3':[val(13),val(14),val(15)],
                      'H_sp2':[val(23),val(24),val(25)],'H_sp3':[val(28),val(29),val(30)]},
          'scaled':{'C':[val(19),val(20)],'H':[val(34),val(35)]}}
    w = openpyxl.load_workbook(asset/'STS.xlsx')
    sts = {}
    num = r'[-+]?\d+(?:\.\d*)?'
    for s in w:
        scale = {t:[float(a),float(b)] for t,a,b in re.findall(
            r'\$D9="([^"]+)",('+num+r')\*H9('+num+r')',s['I9'].value)}
        error = {t:[float(a),float(b)] for t,a,b in re.findall(
            r'\$D9="([^"]+)",2\*\(1-_xlfn.T.DIST\(K9/('+num+r'),('+num+r'),1\)\)',s['L9'].value)}
        heavy = {}
        for condition,a,b in re.findall(r'AND\((.*?)\),('+num+r')\*G9('+num+r')',s['H9'].value):
            kinds = re.findall(r'\$D9="([^"]+)"',condition)
            group = re.findall(r'\$C9="([^"]+)"',condition)
            if len(group) != 1:
                raise ValueError('Unexpected STS heavy atom formula '+condition)
            for kind in kinds:
                heavy[kind+'|'+group[0]] = [float(a),float(b)]
        if len(scale) != 10 or len(error) != 10 or len(heavy) != 20:
            raise ValueError('Upstream STS formula layout changed: '+s.title)
        sts[s.title] = {'scale':scale,'error':error,'heavy':heavy}
    params = {'source':'https://github.com/chemicalchange/EasyCalcNMR_for_OrgSynLab',
      'retrieved':'2026-10-10','workbook_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest()
        for p in [asset/'custom-dp4plus.xlsx',asset/'STS.xlsx']},'dp4':dp,'sts':sts,
      'notes':['Readable DP4 workbook uses Excel default transparent decryption.',
               'Upstream Main!J6 is hardcoded; export restores sigma from selected Statistics column. Native calculator uses table parameters.']}
    source = Path(args.dp4_source)
    constants = {}
    for node in ast.parse(source.read_text(encoding='utf-8')).body:
        if isinstance(node,ast.Assign) and len(node.targets)==1 and isinstance(node.targets[0],ast.Name):
            name = node.targets[0].id
            if name in ['meanC','meanH','stdevC','stdevH']:
                constants[name] = ast.literal_eval(node.value)
    if len(constants)!=4:
        raise ValueError('Classic DP4 parameter layout changed')
    params['dp4_classic'] = {'Goodman_default':{
        'mean':{'C':constants['meanC'],'H':constants['meanH']},
        'sigma':{'C':constants['stdevC'],'H':constants['stdevH']},
        'source':'https://github.com/Goodman-lab/DP5/blob/b6cf559007a5d13fe79654f37daf945ee1661a23/DP4.py',
        'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest()}}
    (root/'calibration.json').write_text(json.dumps(params,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print('Extracted {} DP4+ parameter sets and {} STS solvent tables'.format(len(dp),len(sts)))


if __name__ == '__main__':
    main()
