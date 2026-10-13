"""FC-only H-H exports and explicitly exploratory dJ-DP4 comparisons."""
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import tempfile
import time
from module_options import analysis

# Audited default cells from the authors' calculator bundled by EasyCalcNMR.
# Excel TDIST truncates its degrees-of-freedom argument to an integer.
PARAMETERS={'TMS':{'H':31.830572727272724,'C':192.29325},
            'sigma':{'H':.185,'C':2.306,'J':.992},'df':{'H':14.18,'C':11.38,'J':3.06},
            'j_slope':.9509,'j_intercept':-.1405,
            'source':'https://doi.org/10.1021/acs.orglett.9b01193',
            'workbook_sha256':'8dcb26b2c67fd66e1a11012ef67b078803655c75784dfb93e5342b2d3a2ff275',
            'df_policy':'Excel legacy TDIST integer truncation'}

# Bally & Rablen, JOC 2011: recommended H-only decontraction/core augmentation,
# FC terms multiplied by 0.9155 (Figure 4/probe-set protocol). This is NOT a
# dedicated 2J fit or a calibrated correction for our ordinary 6-31G(d,p) jobs.
BALLY_REFERENCE={'factor':.9155,'source':'https://doi.org/10.1021/jo200513q',
                 'source_protocol':'B3LYP/6-31G(d,p)u+1s[H], FC-only, gas phase',
                 'scope':'general_HH_reference_cross_basis_extrapolation_not_2J_specific'}


def parse_fc(text,natoms):
    """Read one blockwise triangular FC J matrix in Hz; never reduced K or total J."""
    from nmrflow import require,NUMBER
    marker='Fermi Contact (FC) contribution to J (Hz):'
    require(text.count(marker)==1,'Expected exactly one FC contribution to J (Hz) table')
    values={};columns=[]
    for line in text.split(marker,1)[1].splitlines():
        p=line.split()
        if not p:continue
        if all(re.fullmatch(r'\d+',v) for v in p):
            columns=[int(v) for v in p]
            require(columns==sorted(set(columns)) and all(1<=v<=natoms for v in columns),'Invalid FC column header')
            continue
        if columns and re.fullmatch(r'\d+',p[0]) and all(re.fullmatch(NUMBER,v) for v in p[1:]):
            row=int(p[0]);cols=[c for c in columns if c<=row]
            require(1<=row<=natoms and len(p)-1==len(cols),'Truncated/invalid FC matrix row')
            for col,token in zip(cols,p[1:]):
                key=tuple(sorted((row,col)));value=float(token.replace('D','E').replace('d','e'))
                require(key not in values and math.isfinite(value),'Duplicate/nonfinite FC matrix value')
                values[key]=value
            continue
        if values:break
    require(all((a,b) in values for a in range(1,natoms+1) for b in range(a+1,natoms+1)),
            'Incomplete FC J matrix; do not average partial output')
    return {'{},{}'.format(a,b):v for (a,b),v in values.items()}


def pair_key(a,b):return '{},{}'.format(*sorted((a,b)))


def group_average(matrix,atoms_a,atoms_b):
    from nmrflow import require
    require(atoms_a and atoms_b and len(set(atoms_a))==len(atoms_a) and len(set(atoms_b))==len(atoms_b)
            and not set(atoms_a)&set(atoms_b),'Invalid/overlapping J equivalence groups')
    pairs=[pair_key(a,b) for a in atoms_a for b in atoms_b]
    require(all(k in matrix for k in pairs),'Missing assigned J pair')
    # Average signed physical couplings before taking an optional absolute value.
    return sum(matrix[k] for k in pairs)/len(pairs)


def graph_from_mol2(path):
    from nmrflow import require
    text=Path(path).read_text()
    require('@<TRIPOS>BOND' in text,'Missing MOL2 bond connectivity')
    graph={}
    for line in text.split('@<TRIPOS>BOND',1)[1].split('@<TRIPOS>',1)[0].splitlines():
        p=line.split()
        if not p:continue
        a,b=int(p[1]),int(p[2]);graph.setdefault(a,set()).add(b);graph.setdefault(b,set()).add(a)
    return graph


def distance(graph,start,end):
    seen={start};layer={start}
    for d in range(1,5):
        layer={n for a in layer for n in graph.get(a,())}-seen
        if end in layer:return d
        seen.update(layer)
    return None


def worksheet_title(prefix,identifier):
    """Keep short names readable; retain long/case-distinct IDs without collisions."""
    # Reserve room for both longer reference prefixes and a deterministic suffix.
    if len(identifier)<=23:return prefix+identifier
    suffix=hashlib.sha256(identifier.encode('utf-8')).hexdigest()[:10]
    return prefix+identifier[:31-len(prefix)-11]+'_'+suffix


def add_scaled_j_sheet(book,raw_sheet,scheme='dj_dp4'):
    """Add a numeric reference sheet; keep raw cells and probability inputs untouched."""
    from nmrflow import require
    from openpyxl.styles import Alignment,Font,PatternFill
    require(scheme in ('dj_dp4','bally_reference'),'Unknown J worksheet scaling scheme')
    bally=scheme=='bally_reference'
    headers=[cell.value for cell in raw_sheet[1]]
    require(headers and headers[0]=='H atom','Unexpected raw J matrix header')
    atoms=headers[1:]
    require(atoms and all(isinstance(a,int) for a in atoms) and len(set(atoms))==len(atoms),'Invalid J matrix atoms')
    require(raw_sheet.max_row==len(atoms)+1,'Invalid J matrix dimensions')
    title=worksheet_title('Bally参考_' if bally else '经验标度_',raw_sheet.title[2:])
    require(title not in book.sheetnames,'Scaled J sheet already exists: '+title)
    sheet=book.create_sheet(title)
    notes=[
        '整页数值标度后的H–H矩阵（Hz）；原始矩阵见 '+raw_sheet.title,
        '公式：J_scaled = (J_raw - b) / m = (J_raw + 0.1405) / 0.9509',
        '参数：m={}；b={} Hz；来源：{}'.format(PARAMETERS['j_slope'],PARAMETERS['j_intercept'],PARAMETERS['source']),
        '重要：本页将全部非对角元代入同一公式；该经验标度仅适用于³JHH，请自行按原子归属挑选。',
        '默认校准针对B3LYP/6-31G(d,p)//MMFF；本流程DFT几何/自由能权重有差异，标度值仅供参考。',
        '保留原始符号。等价氢组应先平均原始J，再按实验符号约定标度；不要把本页值再次标度或当成实验误差。',
        '²JHH及长程J的本页数值仅是代入公式的结果，未经过相应校准验证；对角线留空。',
    ]
    if bally:
        notes=[
            'Bally–Rablen H–H参考标度（Hz）；从原始矩阵 '+raw_sheet.title+' 独立生成。',
            '公式：J_reference = 0.9155 × J_raw（乘法，无截距；保留正负号）。',
            '来源：'+BALLY_REFERENCE['source']+'；推荐H-only协议/Figure 4。',
            '文献协议：B3LYP/6-31G(d,p)u+1s[H]，气相，FC-only；u为去收缩，+1s为增加紧s函数。',
            '当前原始J使用普通6-31G(d,p)：本页为跨基组外推参考，未按本体系验证，不代表实际使用了+1s。',
            '整张H–H矩阵数值标度，含²J/³J/长程J，由使用者自行挑选；这是总体H–H标度，不是²J专属拟合。',
            '不自动用于dJ-DP4概率，不作第二次标度；等价氢先平均原始带符号J。',
        ]
    last_column=len(atoms)+1
    for row,text in enumerate(notes,1):
        sheet.cell(row,1,text);sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last_column)
        sheet.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center');sheet.row_dimensions[row].height=32
    header_row=9
    for col,value in enumerate(headers,1):sheet.cell(header_row,col,value)
    included=0
    for i,a in enumerate(atoms,2):
        require(raw_sheet.cell(i,1).value==a,'Raw J row/column mismatch')
        sheet.cell(header_row+i-1,1,a)
        for j,b in enumerate(atoms,2):
            if a==b:continue
            raw=raw_sheet.cell(i,j).value
            require(isinstance(raw,(int,float)) and math.isfinite(raw),'Invalid raw J value')
            value=raw*BALLY_REFERENCE['factor'] if bally else (raw-PARAMETERS['j_intercept'])/PARAMETERS['j_slope']
            cell=sheet.cell(header_row+i-1,j,value)
            cell.number_format='0.0000'
            if a<b:included+=1
    for cell in sheet[header_row]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='174A67')
    sheet.freeze_panes='B10';sheet.column_dimensions['A'].width=12
    for column in range(2,last_column+1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(column)].width=12
    return {'sheet':sheet.title,'raw_sheet':raw_sheet.title,'scope':BALLY_REFERENCE['scope'] if bally else 'all_HH_pairs_numerically_scaled_calibration_valid_for_3JHH_only',
            'unique_pairs_scaled':included,'selection':'manual_by_user_no_connectivity_filter',
            'formula':'J_raw * factor' if bally else '(J_raw - j_intercept) / j_slope',
            'parameters':dict(BALLY_REFERENCE) if bally else {k:PARAMETERS[k] for k in ('j_slope','j_intercept')},
            'status':'reference_only_custom_geometry_population','values':'numeric_once_scaled_signed_pairs'}


def add_bally_j_sheet(book,raw_sheet):
    return add_scaled_j_sheet(book,raw_sheet,scheme='bally_reference')


def probabilities(data):
    """Authors' scaled C/H plus scaled J likelihoods, evaluated in log space."""
    from nmrflow import require
    from postprocess import linear_fit,student_tail,normalize
    require(len(data)>=2,'dJ-DP4 relative comparison requires at least two candidates')
    likelihoods={};details=[]
    for cid,rows in data.items():
        sums={}
        for nucleus in ('C','H','J'):
            subset=[r for r in rows if r['type']==nucleus]
            if not subset:continue
            if nucleus=='J':slope,intercept=PARAMETERS['j_slope'],PARAMETERS['j_intercept']
            else:
                x=[r['experimental'] for r in subset]
                y=[PARAMETERS['TMS'][nucleus]-r['calculated'] for r in subset]
                slope,intercept=linear_fit(x,y)
            total=0.0
            for r in subset:
                unscaled=r['calculated'] if nucleus=='J' else PARAMETERS['TMS'][nucleus]-r['calculated']
                scaled=(unscaled-intercept)/slope;error=scaled-r['experimental']
                tail=student_tail(error/PARAMETERS['sigma'][nucleus],int(PARAMETERS['df'][nucleus]))
                require(tail>0,'dJ-DP4 tail underflow; review data and calibration')
                total+=math.log(tail)
                details.append(dict(r,candidate=cid,scaled=scaled,error=error,tail=tail,slope=slope,intercept=intercept))
            sums[nucleus]=total
        require('J' in sums and ('C' in sums or 'H' in sums),'dJ-DP4 requires J and at least one shift nucleus')
        sums['C+H']=sum(sums[k] for k in ('C','H') if k in sums)
        sums['all']=sums['C+H']+sums['J'];likelihoods[cid]=sums
    groups=set.intersection(*(set(v) for v in likelihoods.values()))
    return {g:normalize({c:v[g] for c,v in likelihoods.items()}) for g in sorted(groups)},details


def export_all(cfg,output,config_dir):
    from nmrflow import require,save,digest,gaussian_read,rms,calculation_config
    from postprocess import write_csv,table
    import openpyxl
    output,config_dir=Path(output),Path(config_dir);settings=analysis(cfg)
    parent=output/'module_exports';parent.mkdir(exist_ok=True)
    pending=Path(tempfile.mkdtemp(prefix='.coupling_pending_',dir=str(parent)))
    matrices,shifts,geometries,provenance={},{},{},[];raw=[];weights=[]
    for candidate in cfg['candidates']:
        cid=candidate['id'];manifest=output/cid/'conformers.json';state=json.loads(manifest.read_text())
        require(calculation_config(state['config'])==calculation_config(cfg),'J manifest/config mismatch')
        chosen=[r for r in state['records'] if r['id'] in state['nmr_ids']]
        require(chosen and abs(sum(r['nmr_weight'] for r in chosen)-1)<1e-8,'Invalid J ensemble weights')
        matrices[cid]={};shifts[cid]={};geometries[cid]=chosen[0]['geometry']
        for row in chosen:
            parsed={}
            for stage in ('coupling','coupling_shielding'):
                log=Path(row['coupling'][stage]);receipt=json.loads((log.parent/'receipt.json').read_text())
                require(receipt['output_sha256']==digest(log) and receipt['input_sha256']==digest(log.with_suffix('.gjf')) and receipt['stage']==stage,'J receipt/input/output hash mismatch')
                parsed[stage]=gaussian_read(log,stage)
                require(rms(parsed[stage]['geometry'],row['geometry'])<1e-4,'J/manifest geometry mismatch')
                provenance.append({'candidate':cid,'conformer':row['id'],'stage':stage,'sha256':digest(log),'path':str(log)})
            for pair,value in parsed['coupling']['fc_hz'].items():
                a,b=map(int,pair.split(','))
                if a==b or row['geometry'][a-1][0]!='H' or row['geometry'][b-1][0]!='H':continue
                matrices[cid][pair]=matrices[cid].get(pair,0)+row['nmr_weight']*value
                raw.append({'candidate':cid,'conformer':row['id'],'atom_a':a,'atom_b':b,'fc_hz':value,'weight':row['nmr_weight']})
            for atom,(element,value) in parsed['coupling_shielding']['shielding'].items():
                shifts[cid][atom]=(element,shifts[cid].get(atom,(element,0))[1]+row['nmr_weight']*value)
            weights.append({'candidate':cid,'conformer':row['id'],'full_weight':row['weight'],'property_weight':row['nmr_weight']})
        h=[i+1 for i,a in enumerate(geometries[cid]) if a[0]=='H']
        with (pending/(cid+'_J_result.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f);w.writerow(['H atom']+h)
            for a in h:w.writerow([a]+['X' if a==b else matrices[cid][pair_key(a,b)] for b in h])
    require(raw,'No H-H pairs available for coupling export')
    write_csv(pending/'per_conformer_J.csv',list(raw[0]),raw)
    write_csv(pending/'weights.csv',list(weights[0]),weights)
    averaged=[{'candidate':c,'atom':a,'element':el,'shielding_ppm':v} for c,atoms in shifts.items() for a,(el,v) in atoms.items() if el in ('C','H')]
    write_csv(pending/'weighted_dJ_shielding.csv',list(averaged[0]),averaged)
    status={'status':'raw_complete_experiment_or_assignment_required','scientific_status':'custom_geometry_population_exploratory',
            'protocol_differences':['Gas-phase DFT geometry instead of MMFF','ORCA/Shermo free-energy populations instead of original DFT electronic-energy weights'],
            'parameters':PARAMETERS,'provenance':provenance}
    book=openpyxl.Workbook();main=book.active;main.title='Main';main.append(['Type','Signal','Experimental']+[c['id'] for c in cfg['candidates']])
    expj=config_dir/settings['j_experiment'];mapj=config_dir/settings['j_assignments']
    exps=config_dir/cfg['experiment'];maps=config_dir/cfg['assignments']
    if all(p.is_file() for p in (expj,mapj,exps,maps)):
        ej,es=table(expj),table(exps);mj,ms=table(mapj),table(maps)
        require(ej and es,'Empty J/shift experimental tables')
        require(len({r['signal'] for r in ej})==len(ej) and len({r['signal'] for r in es})==len(es),'Duplicate experimental signal')
        lj={(r['candidate'],r['signal']):r for r in mj};ls={(r['candidate'],r['signal']):r for r in ms}
        require(len(lj)==len(mj) and len(ls)==len(ms),'Duplicate atom mapping')
        data={}
        for cid in matrices:
            data[cid]=[];used_atoms=set();used_pairs=set()
            mol2=output/cid/'sampling/molecule.mol2'
            require(mol2.is_file(),'3J connectivity verification requires the sampled molecule.mol2; XYZ J comparison is not supported')
            graph=graph_from_mol2(mol2)
            for e in es:
                require(e['nucleus'] in ('C','H') and (cid,e['signal']) in ls,'Missing shift atom assignment')
                atoms=[int(a) for a in ls[(cid,e['signal'])]['atoms'].split(';')]
                require(atoms and len(set(atoms))==len(atoms) and not used_atoms&set(atoms),'Duplicate shift atom assignment')
                require(all(a in shifts[cid] and shifts[cid][a][0]==e['nucleus'] for a in atoms),'Shift element/index mismatch')
                used_atoms.update(atoms);value=sum(shifts[cid][a][1] for a in atoms)/len(atoms)
                data[cid].append({'type':e['nucleus'],'signal':e['signal'],'experimental':float(e['ppm']),'calculated':value})
            for e in ej:
                require((cid,e['signal']) in lj and int(e['bond_count'])==3,'Missing J assignment or not 3JHH')
                r=lj[(cid,e['signal'])];a=[int(v) for v in r['atoms_a'].split(';')];b=[int(v) for v in r['atoms_b'].split(';')]
                require(all(0<v<=len(geometries[cid]) and geometries[cid][v-1][0]=='H' for v in a+b),'J assignments must identify H calculation atoms')
                require(all(distance(graph,x,y)==3 for x in a for y in b),'Assigned J pair is not a vicinal 3JHH bond path')
                pairs={pair_key(x,y) for x in a for y in b}
                require(not used_pairs&pairs,'One J pair used by multiple experimental signals');used_pairs.update(pairs)
                value=group_average(matrices[cid],a,b)
                require(e['mode'] in ('signed','magnitude'),'J mode must be signed or magnitude')
                if e['mode']=='magnitude':value=abs(value)
                experimental=float(e['j_hz'])
                require(e['mode']!='magnitude' or experimental>=0,'Magnitude J experiment must be nonnegative')
                data[cid].append({'type':'J','signal':e['signal'],'experimental':experimental,'calculated':value})
            require(all(math.isfinite(r['experimental']) and math.isfinite(r['calculated']) for r in data[cid]),'Nonfinite experimental/calculated value')
        first=next(iter(data.values()));rows=[]
        for i,r in enumerate(first):
            row={'type':r['type'],'signal':r['signal'],'experimental':r['experimental'],**{c:d[i]['calculated'] for c,d in data.items()}}
            rows.append(row);main.append(list(row.values()))
        write_csv(pending/'dJ_DP4_input.csv',list(rows[0]),rows)
        # Export remains possible when the default calibrated geometry does not match.
        status['status']='input_table_complete_exploratory_disabled'
        if settings['coupling_allow_exploratory']:
            result,details=probabilities(data);status.update(status='exploratory_comparison_complete',probabilities_percent=result)
            write_csv(pending/'dJ_DP4_errors.csv',list(details[0]),details)
            sheet=book.create_sheet('Probabilities');sheet.append(['Data']+list(data))
            for g,v in result.items():sheet.append([g]+[v[c] for c in data])
    scaled_sheets=[]
    for cid,matrix in matrices.items():
        sheet=book.create_sheet(worksheet_title('J_',cid));h=[i+1 for i,a in enumerate(geometries[cid]) if a[0]=='H']
        sheet.append(['H atom']+h)
        for a in h:sheet.append([a]+[None if a==b else matrix[pair_key(a,b)] for b in h])
        scaled_sheets.append(add_scaled_j_sheet(book,sheet))
        scaled_sheets.append(add_bally_j_sheet(book,sheet))
    status['scaled_j_worksheets']=scaled_sheets
    sheet=book.create_sheet('Protocol_metadata')
    for k,v in status.items():
        if k!='provenance':sheet.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v)])
    sheet=book.create_sheet('Source_provenance');sheet.append(['candidate','conformer','stage','sha256','path'])
    for r in provenance:sheet.append(list(r.values()))
    book.save(pending/'dJ_DP4_tables.xlsx');save(pending/'comparison_status.json',status)
    final=parent/'coupling'
    if final.exists():final.rename(parent/('coupling_previous_'+time.strftime('%Y%m%d_%H%M%S')+'_'+pending.name[-6:]))
    pending.rename(final)
