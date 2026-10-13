"""Workbook output regression: real numeric scaling, valid scope and raw preservation."""
import tempfile
from pathlib import Path
import unittest
import openpyxl
import coupling as jc


class ScaledSheetChecks(unittest.TestCase):
    def fixture(self):
        book=openpyxl.Workbook();raw=book.active;raw.title='J_demo'
        raw.append(['H atom',1,4,5])
        raw.append([1,None,8.0,-12.0]);raw.append([4,8.0,None,-0.2]);raw.append([5,-12.0,-0.2,None])
        # H1-C2-C3-H4 is vicinal; H1-C2-H5 is geminal. H4-C3-C2-H5 is also vicinal.
        graph={1:{2},2:{1,3,5},3:{2,4},4:{3},5:{2}}
        return book,raw,graph

    def test_full_matrix_signed_scaling_raw_unchanged_and_saved_values(self):
        book,raw,graph=self.fixture();before=list(raw.values)
        info=jc.add_scaled_j_sheet(book,raw);sheet=book[info['sheet']]
        self.assertEqual(before,list(raw.values));self.assertEqual(info['unique_pairs_scaled'],3)
        self.assertAlmostEqual(sheet['C10'].value,(8+.1405)/.9509)
        self.assertAlmostEqual(sheet['B11'].value,sheet['C10'].value)
        self.assertAlmostEqual(sheet['D11'].value,(-.2+.1405)/.9509)
        self.assertAlmostEqual(sheet['D10'].value,(-12+.1405)/.9509);self.assertIsNone(sheet['B10'].value)
        self.assertIn('仅适用于³JHH',sheet['A4'].value)
        self.assertIn('仅供参考',sheet['A5'].value)
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'J.xlsx';book.save(p)
            saved=openpyxl.load_workbook(p,data_only=True)
            self.assertAlmostEqual(saved[info['sheet']]['C10'].value,(8+.1405)/.9509)
            self.assertEqual(before,list(saved['J_demo'].values));saved.close()

    def test_no_connectivity_needed_and_no_double_scaling(self):
        book,raw,_=self.fixture();info=jc.add_scaled_j_sheet(book,raw)
        self.assertEqual(info['unique_pairs_scaled'],3)
        self.assertIn('未经过相应校准验证',book[info['sheet']]['A7'].value)
        with self.assertRaisesRegex(ValueError,'already exists'):jc.add_scaled_j_sheet(book,raw)

    def test_invalid_raw_matrix_is_rejected(self):
        book,raw,_=self.fixture()
        raw['C2']='invalid'
        with self.assertRaisesRegex(ValueError,'Invalid raw J'):jc.add_scaled_j_sheet(book,raw)

    def test_bally_reference_multiplies_original_and_preserves_sign(self):
        book,raw,_=self.fixture();before=list(raw.values)
        dj=jc.add_scaled_j_sheet(book,raw)
        info=jc.add_bally_j_sheet(book,raw);sheet=book[info['sheet']]
        self.assertAlmostEqual(sheet['C10'].value,8*.9155)
        self.assertAlmostEqual(sheet['D10'].value,-12*.9155)
        self.assertAlmostEqual(sheet['B11'].value,sheet['C10'].value)
        self.assertNotAlmostEqual(sheet['C10'].value,book[dj['sheet']]['C10'].value*.9155)
        self.assertEqual(before,list(raw.values));self.assertEqual(info['unique_pairs_scaled'],3)
        self.assertIn('跨基组外推',sheet['A5'].value)
        self.assertIn('不是²J专属拟合',sheet['A6'].value)

    def test_bally_zero_has_no_intercept_and_saved_numeric_value(self):
        book,raw,_=self.fixture();raw['C2']=raw['B3']=0
        info=jc.add_bally_j_sheet(book,raw)
        self.assertEqual(book[info['sheet']]['C10'].value,0)
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'reference.xlsx';book.save(p)
            saved=openpyxl.load_workbook(p,data_only=True)
            self.assertAlmostEqual(saved[info['sheet']]['D10'].value,-12*.9155);saved.close()


if __name__=='__main__':unittest.main()
