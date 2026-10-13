"""Review regressions: queued-input safety, observation and export identity."""
import contextlib
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

import coupling
import nmrflow
import openpyxl
import project_console as pc
import project_views as pv
import postprocess


class ReviewChecks(unittest.TestCase):
    def test_submit_checks_queue_before_sync(self):
        with tempfile.TemporaryDirectory() as d, patch.object(pc,'queue_guard',side_effect=ValueError('active')),patch.object(pc,'sync') as sync:
            with self.assertRaisesRegex(ValueError,'active'):pc.submit(Path(d))
            sync.assert_not_called()

    def test_compare_checks_queue_before_sync(self):
        with patch.object(pc,'queue_guard',side_effect=ValueError('active')),patch.object(pc,'sync') as sync:
            with self.assertRaisesRegex(ValueError,'active'):pc.compare(Path('/unused'))
            sync.assert_not_called()

    def test_underlying_submission_ledger_blocks_duplicate_after_crash(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);(root/'workflow').mkdir()
            (root/'workflow/submissions.tsv').write_text('date\t123.node1\t/job.json\n')
            results=[subprocess.CompletedProcess([],0,'user queue',''),subprocess.CompletedProcess([],0,'job_state = Q','')]
            with patch.object(pc.subprocess,'run',side_effect=results):
                with self.assertRaisesRegex(ValueError,'活动中'):pc.queue_guard(root)

    def test_check_only_does_not_write_project_config_or_preview(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);cfg=json.loads((nmrflow.ROOT/'job.example.json').read_text())
            (root/'job.json').write_text(json.dumps(cfg),encoding='utf-8')
            (root/'project.ini').write_text(pc.settings_text(cfg),encoding='utf-8')
            pc.sync(root)  # Canonical configuration produced by a prepared project.
            before=(root/'job.json').read_bytes()
            preview_before={p.name:p.read_bytes() for p in (root/'参数预览').iterdir()}
            pc.sync(root,write=False)
            self.assertEqual((root/'job.json').read_bytes(),before)
            self.assertEqual(preview_before,{p.name:p.read_bytes() for p in (root/'参数预览').iterdir()})

    def test_live_collector_survives_partial_json(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);cfg=json.loads((nmrflow.ROOT/'job.example.json').read_text())
            cfg.update(output='过程文件',output_layout='direct_v1',candidates=[{'id':'A','input':'inputs/A.pdb'}])
            (root/'job.json').write_text(json.dumps(cfg),encoding='utf-8')
            run=root/'过程文件/A';run.mkdir(parents=True)
            (run/'state.json').write_text('{');(run/'conformers.json').write_text('{')
            result=pv.collect(root)
            self.assertEqual(result['candidates'][0][1],'状态记录暂不可读')
            self.assertTrue((root/'计算总览.html').is_file())

    def test_long_candidate_ids_have_distinct_reference_sheets(self):
        book=openpyxl.Workbook();book.remove(book.active)
        ids=['candidate_'+('long_'*10)+'A','candidate_'+('long_'*10)+'B','a'*23,'a'*24]
        for cid in ids:
            raw=book.create_sheet(coupling.worksheet_title('J_',cid))
            raw.append(['H atom',1,2]);raw.append([1,None,8.0]);raw.append([2,8.0,None])
            coupling.add_scaled_j_sheet(book,raw);coupling.add_bally_j_sheet(book,raw)
        self.assertEqual(len(book.sheetnames),12)
        self.assertTrue(all(len(s)<=31 for s in book.sheetnames))
        self.assertEqual(len({s.casefold() for s in book.sheetnames}),12)

    def test_single_candidate_dj_comparison_is_not_a_probability_validation(self):
        with self.assertRaisesRegex(ValueError,'at least two'):coupling.probabilities({'A':[]})

    def test_nonfinite_thermochemistry_rejected(self):
        for key in ('T','P'):
            cfg=json.loads((nmrflow.ROOT/'job.example.json').read_text())
            cfg['shermo'][key]=float('inf')
            with self.assertRaisesRegex(ValueError,'thermochemistry'):nmrflow.validate_config(cfg)

    def test_nmr_export_rejects_altered_input_and_stage(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);log=root/'job.log';inp=root/'job.gjf'
            log.write_text('accepted Gaussian output');inp.write_text('# NMR input')
            receipt={'stage':'nmr','input_sha256':nmrflow.digest(inp),'output_sha256':nmrflow.digest(log)}
            nmrflow.save(root/'receipt.json',receipt)
            postprocess.verify_nmr_source(log)
            inp.write_text('# modified input')
            with self.assertRaisesRegex(ValueError,'input/output hash'):postprocess.verify_nmr_source(log)
            inp.write_text('# NMR input');receipt['stage']='freq';nmrflow.save(root/'receipt.json',receipt)
            with self.assertRaisesRegex(ValueError,'input/output hash'):postprocess.verify_nmr_source(log)


if __name__=='__main__':unittest.main()
