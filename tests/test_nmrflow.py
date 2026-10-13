"""Scientific/dataflow regressions; no real quantum calculation is run."""
import copy
import json
import math
import os
from pathlib import Path
import sys
import tempfile
import unittest
import re
from unittest.mock import patch

import nmrflow as nf
import postprocess as pp


def gaussian(geom, stage='nmr', energy=-40.0, negative=False):
    text = 'SCF Done: E(RB3LYP) = {} A.U.\n'.format(energy)
    text += ' Standard orientation:\n ----------------------\n Center Atomic Atomic Coordinates\n Number Number Type X Y Z\n ----------------------\n'
    text += '\n'.join('{} {} 0 {} {} {}'.format(i,nf.ELEMENTS.index(a[0]),*a[1:]) for i,a in enumerate(geom,1))
    text += '\n ----------------------\n'
    if stage == 'opt':
        text += 'Optimization completed.\n'
    if stage == 'freq':
        values = [-100 if negative else 100]+[200.0]*(3*len(geom)-7)
        text += 'Frequencies -- '+' '.join(str(v) for v in values)+'\n'
    if stage == 'nmr':
        text += '\n'.join('{} {} Isotropic = {} Anisotropy = 1\n'.format(
            i,a[0],180-i*13 if a[0]=='C' else 33-i*.4) for i,a in enumerate(geom,1))
    return text+'\nNormal termination of Gaussian 16\n'


class ScientificTests(unittest.TestCase):
    def test_ensemble_energy_formats_and_identity(self):
        geom=[['H',0,0,0],['H',0,0,1]]
        blocks=[(' Energy = -54.20 a.u.  #Cluster: 99',geom),
                ('-54.30',geom),('Energy = -5.430D+01 a.u. #Cluster: 1',geom)]
        self.assertEqual([i for i,_,_ in nf.select_ensemble(blocks,2)],[2,3])
        self.assertIs(nf.select_ensemble(blocks,2)[0][2],geom)
        for comment in ['Energy = -54.3 kcal/mol #Cluster: 1','-54.3 eV','Cluster: 1',
                        'Energy = NaN a.u. #Cluster: 1','1.0e999','']:
            with self.assertRaises(ValueError):nf.ensemble_energy(comment)

    def test_sampling_calls_and_stereochemistry(self):
        import sampling as sp
        geom=[['C',0,0,0],['H',1,0,0],['F',0,2,0],['Cl',0,0,3],['Br',-1,-1,-1]]
        reference=[{'center':1,'neighbors':[2,3,4,5],'sign':1 if sp.stereo_volume(geom,[2,3,4,5])>0 else -1}]
        mirror=[[a[0],-a[1],a[2],a[3]] for a in geom]
        kept,rejected=sp.filter_stereochemistry([('good',geom),('bad',mirror)],reference)
        self.assertEqual([c for c,_ in kept],['good'])
        self.assertEqual(rejected[0]['block'],2)
        cfg=json.loads((nf.ROOT/'job.example.json').read_text())
        calls=[]
        with tempfile.TemporaryDirectory() as d:
            work=Path(d)/'sampling'
            source=Path(d)/'molecule.mol2';source.write_text('mock')
            def fake(argv,cwd,log,env=None,stdin=None):
                argv=[str(x) for x in argv];cwd=Path(cwd);calls.append(argv)
                Path(log).write_text('ok')
                if argv[0]==cfg['software']['obabel']: (cwd/'molecule.mol2').write_text('mock')
                elif argv[0].endswith('sobtop'):
                    self.assertTrue(stdin.endswith('4\n\n\n0\n'))
                    for name in ['molecule.top','molecule.itp','molecule.gro']:
                        (cwd/name).write_text('mock')
                elif argv[0]=='gmx' and argv[1]=='mdrun':
                    prefix=argv[argv.index('-deffnm')+1]
                    self.assertEqual(env['OMP_NUM_THREADS'],'4')
                    (cwd/(prefix+'.log')).write_text('Finished mdrun\n')
                elif argv[0]=='gmx' and argv[1]=='check': Path(log).write_text('Last frame 2000 time 200.0\n')
                elif argv[0]=='gmx' and argv[1]=='trjconv': self.assertEqual(stdin,'0\n')
                elif argv[0]==cfg['software']['vmd']:
                    self.assertEqual(env['VMDDIR'],str(Path(cfg['software']['vmd']).parent))
                    self.assertIn('whole.xtc',(cwd/'extract.tcl').read_text())
                    (cwd/'trajectory.xyz').write_text(nf.xyz_text(geom,'-40.0'))
                elif argv[0]==cfg['software']['crest']:
                    Path(log).write_text('CREST terminated normally.')
                    (cwd/'crest_ensemble.xyz').write_bytes(Path(argv[2]).read_bytes())
                elif argv[0].endswith('isostat'): (cwd/'cluster.xyz').write_bytes((cwd/argv[1]).read_bytes())
            with patch.object(sp,'ROOT',nf.ROOT),patch.object(sp,'command',side_effect=fake),patch.object(sp,'resp2',return_value=geom):
                final=sp.sample(source,cfg,work)
                count=len(calls)
                self.assertEqual(sp.sample(source,cfg,work),final)
                self.assertEqual(len(calls),count)
            self.assertEqual(json.loads((work/'sampling.json').read_text())['selected_trajectory'],'md03')

    def test_configuration_and_annealing_guards(self):
        from sampling import md_parameters
        cfg = json.loads((nf.ROOT/'job.example.json').read_text())
        nf.validate_config(cfg)
        for field,value in [('max_conformers',-1),('md_seeds',[1,1,2]),('methods',['01','01'])]:
            bad = copy.deepcopy(cfg)
            bad[field] = value
            with self.assertRaises(ValueError): nf.validate_config(bad)
        bad = copy.deepcopy(cfg)
        bad['postprocessing'][0]['parameter_key']='mPW1PW91|PCM|6-311+G(d,p)'
        with self.assertRaisesRegex(ValueError,'match'): nf.validate_config(bad)
        for i,seed in enumerate(cfg['md_seeds'],1):
            original = (nf.ROOT/'resources'/('md%02d.mdp'%i)).read_text()
            text,duration = md_parameters(original,seed)
            self.assertEqual(duration,200)
            self.assertIn('ld_seed = '+str(seed%2147483647),text)
            with self.assertRaisesRegex(ValueError,'cooling'):
                md_parameters(original.replace('200000','100000'),seed)

    def test_output_lock_and_classification_guards(self):
        with tempfile.TemporaryDirectory() as d:
            with nf.output_lock(Path(d)):
                with self.assertRaisesRegex(ValueError,'Another'):
                    with nf.output_lock(Path(d)): pass
        params = json.loads((nf.ROOT/'calibration.json').read_text())
        rows = [{'nucleus':'C','hybridization':'','carbon_type':'wrong','heavy_atom':'None'}]*3
        for kind in ['dp4plus','sts']:
            branch={'kind':kind,'parameter_key':'CDCl3' if kind=='sts' else 'B3LYP|PCM|6-31G(d)'}
            with self.assertRaises(ValueError): pp.validate_branch_data(branch,{'A':rows},params)

    def test_rejected_frequency_resume(self):
        geom=[['C',0,0,0],['H',1,0,0],['H',0,1,0],['H',0,0,1],['H',-1,0,0]]
        cfg=json.loads((nf.ROOT/'job.example.json').read_text())
        with tempfile.TemporaryDirectory() as d:
            def fake(argv,cwd,log,**kwargs): Path(log).write_text(gaussian(geom,'freq',negative=True))
            with patch.object(nf,'command',side_effect=fake) as solver:
                for _ in range(2):
                    with self.assertRaisesRegex(ValueError,'imaginary'):
                        nf.quantum(nf.ROOT/'resources/04_tem.gjf',geom,cfg,Path(d),'freq')
                self.assertEqual(solver.call_count,1)

    def test_t_distribution_reference(self):
        # Independent analytical Cauchy and Microsoft example reference values.
        for x in [0,.01,.5,1,5,100]:
            self.assertAlmostEqual(pp.student_tail(x,1),.5-math.atan(x)/math.pi,places=12)
        self.assertAlmostEqual(pp.student_tail(1.96,60),.027322463,places=8)
        try:
            from scipy.stats import t
        except ImportError:
            return
        for df in [2,3,5.36432,17,78,292,566]:
            for x in [.0001,.1,1,2,10,100]:
                expected = t.sf(x,df)
                self.assertLess(abs(pp.student_tail(x,df)-expected)/max(expected,1e-200),1e-8)

    def test_parser_rejects_failed_and_imaginary(self):
        geom = [['C',0,0,0],['H',1,0,0],['H',0,1,0]]
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)/'job.log'
            p.write_text(gaussian(geom,'freq',negative=True))
            with self.assertRaisesRegex(ValueError,'imaginary'):
                nf.gaussian_read(p,'freq')
            p.write_text(gaussian(geom).replace('Normal termination of Gaussian','Error termination of Gaussian'))
            with self.assertRaisesRegex(ValueError,'normally'):
                nf.gaussian_read(p,'nmr')
            p.write_text(gaussian(geom).replace('2 H Isotropic','2 C Isotropic'))
            with self.assertRaisesRegex(ValueError,'identity'):
                nf.gaussian_read(p,'nmr')

    def test_rms_preserves_atom_order_and_chirality(self):
        geom = [['C',0,0,0],['H',1,0,0],['F',0,2,0],['Cl',0,0,3]]
        rotated = [[a[0],-a[2]+5,a[1]+7,a[3]-2] for a in geom]
        mirror = [[a[0],-a[1],a[2],a[3]] for a in geom]
        self.assertLess(nf.rms(geom,rotated),1e-10)
        self.assertGreater(nf.rms(geom,mirror),.1)
        with self.assertRaisesRegex(ValueError,'identity'):
            nf.rms(geom,[geom[0],geom[2],geom[1],geom[3]])

    def test_dp4_relative_scores_and_permutation(self):
        params = json.loads((nf.ROOT/'calibration.json').read_text())['dp4']['mPW1PW91|PCM|6-311+G(d,p)']
        rows = [{'nucleus':'C','experimental_ppm':v,'shielding':params['TMS']['C']-(v+.2*((i%2)*2-1)),
                 'hybridization':'sp3' if i<2 else 'sp2'} for i,v in enumerate([15,40,95,130,180])]
        worse = [dict(r,shielding=r['shielding']+e) for r,e in zip(rows,[7,-5,8,-9,12])]
        result = pp.dp4plus({'A':rows,'B':worse},params)['probabilities_percent']
        self.assertGreater(result['all_combined']['A'],99)
        self.assertAlmostEqual(sum(result['all_combined'].values()),100)
        reverse = pp.dp4plus({'A':list(reversed(rows)),'B':list(reversed(worse))},params)['probabilities_percent']
        self.assertAlmostEqual(result['all_combined']['A'],reverse['all_combined']['A'],places=10)
        self.assertNotIn('H_combined',result)

    def test_classic_dp4_goodman_reference(self):
        params = json.loads((nf.ROOT/'calibration.json').read_text())['dp4_classic']['Goodman_default']
        good = [{'nucleus':'C','experimental_ppm':v,'shielding':190-v+e}
                for v,e in zip([15,40,95,130,180],[.2,-.3,.1,-.1,.2])]
        bad = [dict(r,shielding=r['shielding']+e) for r,e in zip(good,[7,-5,8,-9,12])]
        result = pp.dp4({'A':good,'B':bad},params)['probabilities_percent']
        self.assertGreater(result['combined']['A'],99)
        try:
            from scipy import stats
        except ImportError:
            return
        factors = {}
        for cid,rows in [('A',good),('B',bad)]:
            exp = [r['experimental_ppm'] for r in rows]
            shifts = [190-r['shielding'] for r in rows]
            slope,intercept,*_ = stats.linregress(exp,shifts)
            errors = [(s-intercept)/slope-e for s,e in zip(shifts,exp)]
            factors[cid] = math.prod(2*stats.norm.cdf(-abs(e/params['sigma']['C'])) for e in errors)
        self.assertAlmostEqual(result['combined']['A'],100*factors['A']/sum(factors.values()),places=10)

    def test_sts_formula_reference(self):
        params = json.loads((nf.ROOT/'calibration.json').read_text())['sts']['CDCl3']
        exp = [15,40,95,130,180]
        kind = ['sp3-CH3','sp3-CH2/CH/C','sp2-CH2','sp2-CH','C=O']
        rows = []
        for e,k in zip(exp,kind):
            a,b = params['scale'][k]
            rows.append({'nucleus':'C','experimental_ppm':e,'carbon_type':k,'heavy_atom':'None','shielding':(e-b)/a})
        result = pp.sts({'A':rows},params)['A']
        self.assertLess(result['MAE'],1e-12)
        self.assertAlmostEqual(result['Pmean'],1,places=12)
        self.assertEqual(result['Prel_percent'],100)


class PipelineTests(unittest.TestCase):
    @unittest.skipUnless(all((nf.ROOT/'upstream'/name).is_file() for name in ('DP4plus.xlsx','STS.xlsx')),
                         'External upstream workbooks are not installed; see docs/MIGRATION.md')
    def test_weighted_equivalence_and_workbook_cells(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            cfg = json.loads((nf.ROOT/'job.example.json').read_text())
            cfg.update(candidates=[{'id':'candidate_A','input':'unused.xyz'}, {'id':'candidate_B','input':'unused.xyz'}],
                       output=str(root/'runs'))
            for name in ['experiment','assignments']:
                (root/cfg[name]).write_bytes((nf.ROOT/(name+'.example.csv')).read_bytes())
            geom = [['C',0,0,0],['C',1.5,0,0],['C',3,0,0],['H',0,1,0],['H',0,0,1],
                    ['H',0,-1,0],['H',1.5,1,0],['H',3,1,0]]
            for candidate in cfg['candidates']:
                records = []
                for cid,weight,offset in [('c0002',.75,2),('c0001',.25,0)]:
                    logs = {}
                    for method in cfg['methods']:
                        work = root/'runs'/candidate['id']/cid/method
                        work.mkdir(parents=True)
                        text = gaussian(geom)
                        text = re.sub(r'Isotropic = ([-\d.]+)',lambda m:'Isotropic = '+str(float(m[1])+offset),text)
                        log = work/'job.log'
                        log.write_text(text)
                        inp=log.with_suffix('.gjf');inp.write_text('# test NMR input\n')
                        nf.save(work/'receipt.json',{'output_sha256':nf.digest(log),'input_sha256':nf.digest(inp),'stage':'nmr'})
                        logs[method] = str(log)
                    records.append({'id':cid,'geometry':geom,'nmr_weight':weight,'nmr':logs})
                nf.save(root/'runs'/candidate['id']/'conformers.json',{'config':cfg,'records':records,
                        'nmr_ids':['c0001','c0002'],'coverage_within_computed_ensemble':.99,
                        'ensemble_total':8,'dft_selected':6})
            pp.export(cfg,root/'runs',root)
            rows = pp.table(root/'runs/exports/weighted_shielding.csv')
            row = next(r for r in rows if r['candidate']=='candidate_A' and r['method']=='02' and r['atom']=='1')
            self.assertAlmostEqual(float(row['shielding_ppm']),168.5)
            rows = pp.table(root/'runs/exports/method02_comparison_input.csv')
            self.assertAlmostEqual(float(next(r for r in rows if r['signal']=='H1')['candidate_A']),32.5)
            w = openpyxl.load_workbook(root/'runs/exports/dp4plus_method02_exploratory.xlsx')
            self.assertAlmostEqual(w['Main']['D14'].value,168.5)
            self.assertEqual(w['Main']['E6'].value,'6-311+G(d,p)')
            self.assertEqual(w['Main']['J6'].value,.110542)
            meta=dict(w['Protocol_metadata'].values)
            self.assertEqual(meta['NMR level'],'mPW1PW91/6-311+G(d,p)')
            self.assertEqual(json.loads(meta['NMR settings'])['gaussian_defaults'],'G09Defaults')
            w = openpyxl.load_workbook(root/'runs/exports/dp4plus_method01_exploratory.xlsx')
            self.assertEqual(w['Main']['A6'].value,'B3LYP')
            self.assertEqual(w['Main']['C6'].value,'PCM')
            self.assertEqual(w['Main']['E6'].value,'6-31G(d)')
            self.assertAlmostEqual(w['Main']['D14'].value,168.5)
            self.assertFalse((root/'runs/exports/dp4_method01_exploratory.xlsx').exists())
            results = json.loads((root/'runs/exports/comparison_results.json').read_text())
            self.assertIn('dp4plus_01',results['statistics'])
            self.assertIn('dp4plus_02',results['statistics'])
            self.assertNotIn('dp4_01',results['statistics'])
            self.assertEqual(results['method_layout'],'dp4_dp4_sts_v4')
            self.assertEqual(results['nmr_settings']['solvent_model'],'IEFPCM')
            self.assertEqual(results['clustering_profile'],'coarse')
            w = openpyxl.load_workbook(root/'runs/exports/sts_method03_exploratory.xlsx')
            self.assertAlmostEqual(w['CDCl3']['G9'].value,168.5)
            self.assertEqual(w['CDCl3']['D9'].value,'sp3-CH3')
            previous = nf.digest(root/'runs/exports/comparison_results.json')
            text = (root/'assignments.csv').read_text().replace('H1,4;5;6','H1,2')
            (root/'assignments.csv').write_text(text)
            with self.assertRaisesRegex(ValueError,'multiple|element'):
                pp.export(cfg,root/'runs',root)
            self.assertEqual(previous,nf.digest(root/'runs/exports/comparison_results.json'))

    def test_full_pipeline_isolation_resume_exports(self):
        original_root = nf.ROOT
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'resources').mkdir()
            for p in (original_root/'resources').glob('*.gjf'):
                (root/'resources'/p.name).write_bytes(p.read_bytes())
            (root/'resources/05_tem.inp').write_bytes((original_root/'resources/05_tem.inp').read_bytes())
            (root/'resources/shermo_settings.ini').write_text('T= 298.15\nP= 1\nilowfreq= 2\n')
            cfg = json.loads((original_root/'job.example.json').read_text())
            geom = [['C',0,0,0],['H',1,0,0],['H',0,1,0],['H',0,0,1],['H',-1,0,0]]
            source = root/'input.xyz'
            source.write_text(nf.xyz_text(geom,'-40.0'))
            cfg.update(candidates=[{'id':'A','input':str(source)}],output=str(root/'runs'),
                       methods=['01','02','03'],postprocessing=[],output_layout='direct_v1')
            calls = []

            def fake_command(argv,cwd,log,env=None,stdin=None):
                argv = [str(v) for v in argv]
                cwd, log = Path(cwd),Path(log)
                calls.append((argv,str(cwd)))
                if argv[0].endswith('isostat'):
                    (cwd/'cluster.xyz').write_bytes(Path(argv[1]).read_bytes())
                    log.write_text('cluster completed')
                elif argv[0].endswith('/orca'):
                    self.assertTrue(env['PATH'].startswith(cfg['software']['openmpi']))
                    log.write_text('FINAL SINGLE POINT ENERGY -40.1\nORCA TERMINATED NORMALLY\n')
                elif argv[0].endswith('/Shermo'):
                    self.assertEqual(argv[2], '-E')
                    self.assertEqual(len(argv), 4)
                    settings = (cwd/'settings.ini').read_text()
                    self.assertIn('intpvib= 100', settings)
                    self.assertIn('defmass= 3', settings)
                    log.write_text('Sum of electronic energy and thermal correction to G: -40.05\n')
                else:
                    self.assertRegex(stdin,r'\n\n0 1\nC ')
                    self.assertTrue(stdin.endswith('\n\n'))
                    stage = 'opt' if ' opt' in stdin.lower() else ('freq' if '#p freq' in stdin.lower() else 'nmr')
                    log.write_text(gaussian(geom,stage))
            with patch.object(nf,'ROOT',root),patch.object(nf,'command',fake_command):
                nf.run_candidate(cfg['candidates'][0],cfg,root/'runs/A')
                first_calls = len([a for a,_ in calls if a[0].endswith('/g16')])
                nf.run_candidate(cfg['candidates'][0],cfg,root/'runs/A')
                self.assertEqual(first_calls,len([a for a,_ in calls if a[0].endswith('/g16')]))
            pp.export(cfg,root/'runs',root)
            self.assertTrue((root/'runs/exports/weighted_shielding.csv').exists())
            self.assertEqual(json.loads((root/'runs/exports/comparison_status.json').read_text())['status'],
                             'raw_exports_complete_assignment_required')
            result = json.loads((root/'runs/A/conformers.json').read_text())
            self.assertEqual(result['nmr_ids'],['c0001'])
            self.assertEqual(result['records'][0]['weight'],1)
            log = Path(result['records'][0]['nmr']['01'])
            log.write_text(log.read_text()+'tamper')
            with patch.object(nf,'ROOT',root),patch.object(nf,'command',fake_command):
                with self.assertRaisesRegex(ValueError,'Changed input/output'):
                    nf.run_candidate(cfg['candidates'][0],cfg,root/'runs/A')


if __name__ == '__main__':
    unittest.main()
